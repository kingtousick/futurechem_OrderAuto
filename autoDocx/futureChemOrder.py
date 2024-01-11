from PyQt5.QtWidgets import QApplication, QWidget, QLabel, QLineEdit, QPushButton, QVBoxLayout, QFileDialog, QTabWidget, QTableWidget, QTableWidgetItem, QHBoxLayout
from PyQt5.QtCore import Qt
import sys
from docx import Document
from datetime import datetime

class PurchaseRequestApp(QWidget):
    def __init__(self):
        super().__init__()
        
        self.setWindowTitle("구매요구서 자동화")
        self.setFixedSize(300, 930)

        self.tab_widget = QTabWidget(self)
        self.setup_tab1()
        # self.setup_tab2()
        # self.setup_tab3()

        layout = QVBoxLayout()
        layout.addWidget(self.tab_widget)
        self.setLayout(layout)
       
    def setup_tab1(self):
               
        tab1 = QWidget()
        self.tab_widget.addTab(tab1, "수동입력")
        
        label = QLabel("구매요구서 양식:")
        self.file_path = QLineEdit(self)
        self.file_path.setReadOnly(True)
        button = QPushButton("Browse", self)
        button.clicked.connect(self.upload_form)
    
        
        
        labels = ["기관명", "기관코드", "핵종명[앞]", "핵종수량", "핵종명[뒤]", "모델명"]
        manpower_labels = [f"투입인력 : {i:02d}:{j:02d}" for i in range(8, 18) for j in range(0, 60, 30) if not (i == 17 and j == 30)]        
        # 투입인력 항목 추가
        labels +=manpower_labels + ["구입 예정일"]
        
        table = QTableWidget(self)
        table.setColumnCount(2)  # 2개의 열을 갖는 테이블
        table.setRowCount(len(labels))  # 레이블의 수에 따라 테이블 행 수 설정
        table.setHorizontalHeaderLabels(["명칭", "입력값"])
        table.setColumnWidth(1, 150) 
        # 테이블에 레이블과 편집 가능한 QLineEdit 위젯 추가
        for i, labelT in enumerate(labels):
            table.setItem(i, 0, QTableWidgetItem(labelT))
            table.setCellWidget(i, 1, QLineEdit(self))

        # Row 번호 숨기기
        table.verticalHeader().setVisible(False)        

        generate_button = QPushButton("파일생성하기", self)
        generate_button.clicked.connect(self.generate_word_file)

        tab1_layout = QVBoxLayout()
        hbox_file_path = QHBoxLayout()
        hbox_file_path.addWidget(label)
        hbox_file_path.addWidget(self.file_path)
        hbox_file_path.addWidget(button)
        tab1_layout.addLayout(hbox_file_path)
        
        # 테이블 LayOut 등록 
        tab1_layout.addWidget(table)
        
        # 파일생성 버튼 등록 
        tab1_layout.addWidget(generate_button)
   

        tab1.setLayout(tab1_layout) 
    
    # # 텍스트 변경 
    def replace_text_in_cell(self, cell, key, value):
        cell.text = cell.text.replace(f"{key}", value)
    
    # # 값받아오기 변경 
    def get_value_from_table(self, row):
        item = self.tab_widget.widget(0).layout().itemAt(row, 1)
        if item and isinstance(item.widget(), QLineEdit):
            return item.widget().text()
        return ""
    
    # # 파일올리기 
    def upload_form(self):
        file_dialog = QFileDialog()
        file_path = file_dialog.getOpenFileName(self, "Select Word file", "", "Word files (*.docx)")[0]
        if file_path:
            self.file_path.setText(file_path)
            
    def generate_word_file(self):
        template_path = self.file_path.text()
        if template_path:
            document = Document(template_path)
            placeholders = {
                "instNm": self.get_value_from_table(0),
                "instCd": self.get_value_from_table(1),
                "isotopeFront": self.get_value_from_table(2),
                "Quantity": self.get_value_from_table(3),
                "isotopeEnd": self.get_value_from_table(4),
                "model": self.get_value_from_table(5),
                "manPw1": self.get_value_from_table(6),
                "manPw2": self.get_value_from_table(7),
                "manPw3": self.get_value_from_table(8),
                "manPw4": self.get_value_from_table(9),
                "manPw5": self.get_value_from_table(10),
                "manPw6": self.get_value_from_table(11),
                "manPw7": self.get_value_from_table(12),
                "manPw8": self.get_value_from_table(13),
                "manPw9": self.get_value_from_table(14),
                "manPw10": self.get_value_from_table(15),
                "manPw11": self.get_value_from_table(16),
                "manPw12": self.get_value_from_table(17),
                "manPw13": self.get_value_from_table(18),
                "manPw14": self.get_value_from_table(19),
                "manPw15": self.get_value_from_table(20),
                "manPw16": self.get_value_from_table(21),
                "manPw17": self.get_value_from_table(22),
                "manPw18": self.get_value_from_table(23)
            }

            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for key, value in placeholders.items():
                            self.replace_text_in_cell(cell, key, value)

            now = datetime.now()
            timestamp = now.strftime("%Y%m%d_%H%M%S")
            new_file_path = f"구매요구서_{timestamp}.docx"
            document.save(new_file_path)
            print(f"파일이 저장되었습니다: {new_file_path}")
            
    
    def setup_tab2(self):
        tab2 = QWidget()
        self.tab_widget.addTab(tab2, "데이터 등록")
        # Implement Tab 2 UI and functionality here

    def setup_tab3(self):
        tab3 = QWidget()
        self.tab_widget.addTab(tab3, "검색 입력")
        # Implement Tab 3 UI and functionality here        

def main():
    app = QApplication(sys.argv)
    window = PurchaseRequestApp()
    window.show()
    sys.exit(app.exec_())

if __name__ == "__main__":
    main()  
    
    

